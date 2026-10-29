"""把一次仿真运行导出成 Excel 工作簿。

设计要点（每一条都是刻意的）：
1. **只用库里存的结果**：数据源是 simulation_runs.result_json，不重新算一遍。
   否则"导出时配置已改"会导致表格数字与当时页面显示不一致，审计就失去了意义。
2. **不落盘**：用 BytesIO 在内存里生成，直接作为响应体返回。服务器不产生垃圾文件。
3. **拿到就能用**：中文列头、冻结首行、列宽、表头底色。
4. **多 sheet 分维度**：摘要 / 列车 / AGV / 作业 / 风险。
"""

from __future__ import annotations

import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0E7490")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_sheet(
    worksheet,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int],
) -> None:
    """统一的写表逻辑：表头样式 + 冻结首行 + 列宽。"""
    worksheet.title = title
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN

    for row in rows:
        worksheet.append(row)

    worksheet.freeze_panes = "A2"
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _round(value: Any, digits: int = 2) -> Any:
    try:
        return round(float(value), digits)
    except (TypeError, ValueError):
        return value


def build_workbook(run: dict) -> bytes:
    """run 是 simulation_runs 的一行（含 request_json / result_json 两个字符串）。"""
    request = json.loads(run["request_json"])
    result = json.loads(run["result_json"])

    order = request.get("order") or {}
    metrics = result.get("metrics") or {}
    summary = result.get("routeSummary") or {}
    train_plan = result.get("trainPlan") or []
    agv_plan = result.get("agvPlan") or []
    operation_rows = result.get("operationRows") or []
    risk = result.get("riskAssessment") or {}

    workbook = Workbook()

    # ── Sheet 1：运行摘要 ──
    overview_rows: list[list[Any]] = [
        ["运行编号", run.get("id")],
        ["生成时间", run.get("created_at")],
        ["起点 → 终点", f"{run.get('origin')} → {run.get('destination')}"],
        ["货运重量 (kg)", _round(order.get("cargoWeightKg"), 1)],
        ["货运件数 (件)", order.get("cargoCount")],
        ["列车段数", len(train_plan)],
        ["总站数", metrics.get("totalStationCount")],
        ["总时长 (分钟)", metrics.get("totalDurationMin")],
        ["换乘次数", metrics.get("transferCount")],
        ["AGV 作业节点", metrics.get("agvOperationNodeCount")],
        ["接力次数", metrics.get("relayCount")],
        ["收入 (元)", _round(metrics.get("revenue"))],
        ["成本 (元)", _round(metrics.get("totalCost"))],
        ["利润 (元)", _round(metrics.get("profit"))],
        ["碳减排 (kg)", _round(metrics.get("carbonSavedKg"))],
        ["容量利用率", f"{_round((metrics.get('capacityUtilization') or 0) * 100, 1)}%"],
        ["平台交接 (秒)", metrics.get("platformHandoffSeconds")],
        ["AGV 作业 (秒)", metrics.get("agvOperationSeconds")],
    ]
    _write_sheet(workbook.active, "运行摘要", ["项目", "值"], overview_rows, [22, 40])

    # ── Sheet 2：列车编排 ──
    train_rows = [
        [
            index + 1,
            plan.get("trainId"),
            plan.get("lineName"),
            plan.get("direction"),
            plan.get("boardingStation"),
            plan.get("alightingStation"),
            plan.get("stationCount"),
            plan.get("durationMin"),
            plan.get("departureTime"),
            plan.get("arrivalTime"),
        ]
        for index, plan in enumerate(train_plan)
    ]
    _write_sheet(
        workbook.create_sheet(),
        "列车编排",
        ["序号", "车次", "线路", "方向", "上车站", "下车站", "站数", "运行(分)", "发车", "到达"],
        train_rows,
        [6, 24, 20, 14, 12, 12, 8, 10, 10, 10],
    )

    # ── Sheet 3：AGV 编排 ──
    agv_rows = [
        [
            index + 1,
            node.get("agvId"),
            node.get("station"),
            node.get("role"),
            node.get("source"),
            node.get("battery"),
            node.get("capacityKg"),
            node.get("bindTrain"),
            node.get("operationTime"),
            node.get("action"),
        ]
        for index, node in enumerate(agv_plan)
    ]
    _write_sheet(
        workbook.create_sheet(),
        "AGV 编排",
        ["序号", "AGV 编号", "站点", "角色", "分配来源", "电量(%)", "载重(kg)", "绑定车次", "作业时刻", "动作"],
        agv_rows,
        [6, 16, 12, 12, 20, 10, 10, 24, 12, 40],
    )

    # ── Sheet 4：作业明细 ──
    operation_data = [
        [
            index + 1,
            row.get("orderId"),
            row.get("orderType"),
            row.get("agvId"),
            row.get("flightSegment"),
            row.get("bindTrain"),
            row.get("boardingStop"),
            row.get("boardingTime"),
            row.get("AlightingStop"),
            row.get("alightingTime"),
        ]
        for index, row in enumerate(operation_rows)
    ]
    _write_sheet(
        workbook.create_sheet(),
        "作业明细",
        ["序号", "订单编号", "订单类型", "AGV 编号", "航段", "绑定车次", "上车站", "上车时间", "下车站", "下车时间"],
        operation_data,
        [6, 12, 12, 16, 14, 24, 12, 12, 12, 12],
    )

    # ── Sheet 5：风险评估 ──
    risk_rows = [
        [key, item.get("level"), item.get("message")]
        for key, item in risk.items()
        if isinstance(item, dict)
    ]
    _write_sheet(
        workbook.create_sheet(),
        "风险评估",
        ["检查项", "等级", "说明"],
        risk_rows,
        [22, 12, 60],
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
