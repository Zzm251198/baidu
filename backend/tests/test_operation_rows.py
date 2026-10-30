"""作业明细（operationRows）与审计归因的后端测试。

对应本轮修复的两处：
1. 终端节点不再复用上一段列车的上下车信息（否则与「中转换挂」行完全相同）；
2. 仿真记录与 Excel 导出都能归因到「哪种派车算法 + 哪条路径策略」。

这两条都属于"数据看着对、其实无法复盘"的问题，必须由测试钉住。
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.services import simulation


def test_operation_rows_terminal_node_is_not_duplicated(sample_order, sample_route):
    """终端节点：没有"上车"，只有到达；且任意两行的站点+时间不相同。"""
    rows = simulation.simulate(sample_order, sample_route)["operationRows"]

    assert len(rows) == 3
    first, middle, last = rows

    assert first["boardingStop"] == "牛湖"
    assert first["AlightingStop"] == "会展中心"

    assert middle["boardingStop"] == "会展中心"
    assert middle["AlightingStop"] == "罗湖"

    # 终端节点不再复制上一段
    assert last["flightSegment"] == "终端出站"
    assert last["boardingStop"] == "—"
    assert last["boardingTime"] == "—"
    assert last["AlightingStop"] == "罗湖"

    keys = {
        (row["boardingStop"], row["boardingTime"], row["AlightingStop"], row["alightingTime"]) for row in rows
    }
    assert len(keys) == len(rows), "出现两行完全相同 → 表格里有一条是废信息"


def test_operation_rows_parity_with_frontend(sample_order, sample_route):
    """前后端同口径：站点/时间的取值规则必须一致（差一处就会两套表格）。"""
    rows = simulation.simulate(sample_order, sample_route)["operationRows"]

    # 始发/中转沿用列车段，终端用占位符——这条规则前端 simulationEngine.js 必须一致
    assert [row["boardingStop"] for row in rows] == ["牛湖", "会展中心", "—"]
    assert [row["AlightingStop"] for row in rows] == ["会展中心", "罗湖", "罗湖"]


def test_simulations_list_reports_algorithm_and_route_strategy(client, sample_order, sample_route):
    """审计列表要能回答"这条记录是怎么跑出来的"。"""
    result = client.post(
        "/api/simulate",
        json={
            "order": sample_order,
            "route": sample_route,
            "algorithm": "capacity-first",
            "routeStrategy": "faster",
        },
    ).json()
    assert result["success"] is True

    runs = client.get("/api/simulations", params={"limit": 10}).json()
    row = next(item for item in runs if item["id"] == result["runId"])

    assert row["algorithm"] == "capacity-first"
    assert row["routeStrategy"] == "faster"


def test_simulations_list_defaults_for_legacy_payload(client, sample_order, sample_route):
    """不传开关时回落默认（baseline / smart），老记录同样能读出默认值。"""
    result = client.post("/api/simulate", json={"order": sample_order, "route": sample_route}).json()

    runs = client.get("/api/simulations", params={"limit": 10}).json()
    row = next(item for item in runs if item["id"] == result["runId"])

    assert row["algorithm"] == "baseline"
    assert row["routeStrategy"] == "smart"


def test_export_summary_includes_audit_fields(client, sample_order, sample_route):
    """导出的 Excel 也要自证：摘要 sheet 里有调度算法与路径策略两行。"""
    result = client.post(
        "/api/simulate",
        json={
            "order": sample_order,
            "route": sample_route,
            "algorithm": "capacity-first",
            "routeStrategy": "fewer",
        },
    ).json()

    content = client.get(f"/api/simulations/{result['runId']}/export.xlsx").content
    workbook = load_workbook(io.BytesIO(content))
    values = {row[0]: row[1] for row in workbook["运行摘要"].iter_rows(min_row=2, values_only=True)}

    assert values["调度算法"] == "capacity-first"
    assert values["路径策略"] == "fewer"
